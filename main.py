from fastapi import FastAPI, APIRouter, HTTPException, Depends
from fastapi import Response
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict, EmailStr
from typing import List, Optional
import uuid
from datetime import datetime, timezone, timedelta
import jwt
from passlib.context import CryptContext
import io
import httpx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Security
SECRET_KEY = os.environ.get('SECRET_KEY', 'your-secret-key-change-in-production')
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 7  # 7 days

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
security = HTTPBearer()

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")


# Helper functions
def hash_password(password: str) -> str:
    return pwd_context.hash(password)

def verify_password(plain_password: str, hashed_password: str) -> bool:
    return pwd_context.verify(plain_password, hashed_password)

def create_access_token(data: dict) -> str:
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    try:
        token = credentials.credentials
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id: str = payload.get("sub")
        if user_id is None:
            raise HTTPException(status_code=401, detail="Invalid authentication credentials")
        return user_id
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token has expired")
    except jwt.JWTError:
        raise HTTPException(status_code=401, detail="Invalid authentication credentials")

# Define Models
class User(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    email: EmailStr
    name: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class UserCreate(BaseModel):
    email: EmailStr
    name: str
    password: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class LessonPlan(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: str
    subject: str
    grade_level: str
    topic: str
    objectives: str
    materials: str
    procedure: str
    assessment: str
    notes: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class LessonPlanCreate(BaseModel):
    subject: str
    grade_level: str
    topic: str
    objectives: str
    materials: str
    procedure: str
    assessment: str
    notes: Optional[str] = ""

class Activity(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: str
    title: str
    description: str
    activity_type: str
    duration: str
    materials: str
    instructions: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class ActivityCreate(BaseModel):
    title: str
    description: str
    activity_type: str
    duration: str
    materials: str
    instructions: str

class PresentationModel(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: str
    title: str
    subject: str
    slides_content: str
    notes: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class PresentationCreate(BaseModel):
    title: str
    subject: str
    slides_content: str
    notes: Optional[str] = ""

# Routes
@api_router.get("/")
async def root():
    return {"message": "Lesson Planner API"}

# Auth Routes
@api_router.post("/auth/register")
async def register(user_data: UserCreate):
    # Check if user exists
    existing_user = await db.users.find_one({"email": user_data.email})
    if existing_user:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    # Create user
    user = User(
        email=user_data.email,
        name=user_data.name
    )
    
    user_dict = user.model_dump()
    user_dict['created_at'] = user_dict['created_at'].isoformat()
    user_dict['password'] = hash_password(user_data.password)
    
    await db.users.insert_one(user_dict)
    
    # Create token
    access_token = create_access_token(data={"sub": user.id})
    
    return {
        "access_token": access_token,
        "token_type": "bearer",
        "user": {
            "id": user.id,
            "email": user.email,
            "name": user.name
        }
    }

@api_router.post("/auth/login")
async def login(credentials: UserLogin):
    # Find user
    user = await db.users.find_one({"email": credentials.email})
    if not user:
        raise HTTPException(status_code=401, detail="Invalid email or password")
    
    # Verify password
    if not verify_password(credentials.password, user['password']):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    
    # Create token
    access_token = create_access_token(data={"sub": user['id']})
    
    return {
        "access_token": access_token,
        "token_type": "bearer",
        "user": {
            "id": user['id'],
            "email": user['email'],
            "name": user['name']
        }
    }

@api_router.get("/auth/me")
async def get_me(user_id: str = Depends(get_current_user)):
    user = await db.users.find_one({"id": user_id}, {"_id": 0, "password": 0})
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user

# Lesson Plans
@api_router.post("/lesson-plans", response_model=LessonPlan)
async def create_lesson_plan(input: LessonPlanCreate, user_id: str = Depends(get_current_user)):
    lesson_dict = input.model_dump()
    lesson_dict['user_id'] = user_id
    lesson_obj = LessonPlan(**lesson_dict)
    
    doc = lesson_obj.model_dump()
    doc['created_at'] = doc['created_at'].isoformat()
    
    await db.lesson_plans.insert_one(doc)
    return lesson_obj

@api_router.get("/lesson-plans", response_model=List[LessonPlan])
async def get_lesson_plans(user_id: str = Depends(get_current_user)):
    plans = await db.lesson_plans.find({"user_id": user_id}, {"_id": 0}).to_list(1000)
    
    for plan in plans:
        if isinstance(plan['created_at'], str):
            plan['created_at'] = datetime.fromisoformat(plan['created_at'])
    
    return plans

@api_router.delete("/lesson-plans/{plan_id}")
async def delete_lesson_plan(plan_id: str, user_id: str = Depends(get_current_user)):
    result = await db.lesson_plans.delete_one({"id": plan_id, "user_id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Lesson plan not found")
    return {"message": "Deleted successfully"}

# Activities
@api_router.post("/activities", response_model=Activity)
async def create_activity(input: ActivityCreate, user_id: str = Depends(get_current_user)):
    activity_dict = input.model_dump()
    activity_dict['user_id'] = user_id
    activity_obj = Activity(**activity_dict)
    
    doc = activity_obj.model_dump()
    doc['created_at'] = doc['created_at'].isoformat()
    
    await db.activities.insert_one(doc)
    return activity_obj

@api_router.get("/activities", response_model=List[Activity])
async def get_activities(user_id: str = Depends(get_current_user)):
    activities = await db.activities.find({"user_id": user_id}, {"_id": 0}).to_list(1000)
    
    for activity in activities:
        if isinstance(activity['created_at'], str):
            activity['created_at'] = datetime.fromisoformat(activity['created_at'])
    
    return activities

@api_router.delete("/activities/{activity_id}")
async def delete_activity(activity_id: str, user_id: str = Depends(get_current_user)):
    result = await db.activities.delete_one({"id": activity_id, "user_id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Activity not found")
    return {"message": "Deleted successfully"}

# Presentations
@api_router.post("/presentations", response_model=PresentationModel)
async def create_presentation(input: PresentationCreate, user_id: str = Depends(get_current_user)):
    presentation_dict = input.model_dump()
    presentation_dict['user_id'] = user_id
    presentation_obj = PresentationModel(**presentation_dict)
    
    doc = presentation_obj.model_dump()
    doc['created_at'] = doc['created_at'].isoformat()
    
    await db.presentations.insert_one(doc)
    return presentation_obj

@api_router.get("/presentations", response_model=List[PresentationModel])
async def get_presentations(user_id: str = Depends(get_current_user)):
    presentations = await db.presentations.find({"user_id": user_id}, {"_id": 0}).to_list(1000)
    
    for presentation in presentations:
        if isinstance(presentation['created_at'], str):
            presentation['created_at'] = datetime.fromisoformat(presentation['created_at'])
    
    return presentations

@api_router.delete("/presentations/{presentation_id}")
async def delete_presentation(presentation_id: str, user_id: str = Depends(get_current_user)):
    result = await db.presentations.delete_one({"id": presentation_id, "user_id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Presentation not found")
    return {"message": "Deleted successfully"}

# Download endpoints for Lesson Plans
@api_router.get("/lesson-plans/{plan_id}/download/word")
async def download_lesson_plan_word(plan_id: str, user_id: str = Depends(get_current_user)):
    plan = await db.lesson_plans.find_one({"id": plan_id, "user_id": user_id}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=404, detail="Lesson plan not found")
    
    # Create Word document
    doc = Document()
    
    # Title
    title = doc.add_heading('Lesson Plan', 0)
    title.alignment = 1  # Center
    
    # Add content
    doc.add_heading(f"{plan['subject']} - {plan['topic']}", level=1)
    
    doc.add_paragraph(f"Grade Level: {plan['grade_level']}")
    doc.add_paragraph()
    
    doc.add_heading('Learning Objectives', level=2)
    doc.add_paragraph(plan['objectives'])
    
    doc.add_heading('Materials Needed', level=2)
    doc.add_paragraph(plan['materials'])
    
    doc.add_heading('Procedure/Activities', level=2)
    doc.add_paragraph(plan['procedure'])
    
    doc.add_heading('Assessment Methods', level=2)
    doc.add_paragraph(plan['assessment'])
    
    if plan.get('notes'):
        doc.add_heading('Additional Notes', level=2)
        doc.add_paragraph(plan['notes'])
    
    # Save to BytesIO
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    filename = f"{plan['subject']}_{plan['topic']}_lesson_plan.docx".replace(' ', '_')
    
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@api_router.get("/lesson-plans/{plan_id}/download/pdf")
async def download_lesson_plan_pdf(plan_id: str, user_id: str = Depends(get_current_user)):
    plan = await db.lesson_plans.find_one({"id": plan_id, "user_id": user_id}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=404, detail="Lesson plan not found")
    
    # Create PDF
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    
    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=RGBColor(0, 0, 128),
        spaceAfter=30,
        alignment=1
    )
    story.append(Paragraph('Lesson Plan', title_style))
    story.append(Spacer(1, 0.3*inch))
    
    # Subject and Topic
    story.append(Paragraph(f"<b>{plan['subject']} - {plan['topic']}</b>", styles['Heading2']))
    story.append(Paragraph(f"Grade Level: {plan['grade_level']}", styles['Normal']))
    story.append(Spacer(1, 0.2*inch))
    
    # Sections
    story.append(Paragraph('<b>Learning Objectives</b>', styles['Heading3']))
    story.append(Paragraph(plan['objectives'], styles['Normal']))
    story.append(Spacer(1, 0.2*inch))
    
    story.append(Paragraph('<b>Materials Needed</b>', styles['Heading3']))
    story.append(Paragraph(plan['materials'], styles['Normal']))
    story.append(Spacer(1, 0.2*inch))
    
    story.append(Paragraph('<b>Procedure/Activities</b>', styles['Heading3']))
    story.append(Paragraph(plan['procedure'], styles['Normal']))
    story.append(Spacer(1, 0.2*inch))
    
    story.append(Paragraph('<b>Assessment Methods</b>', styles['Heading3']))
    story.append(Paragraph(plan['assessment'], styles['Normal']))
    
    if plan.get('notes'):
        story.append(Spacer(1, 0.2*inch))
        story.append(Paragraph('<b>Additional Notes</b>', styles['Heading3']))
        story.append(Paragraph(plan['notes'], styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    
    filename = f"{plan['subject']}_{plan['topic']}_lesson_plan.pdf".replace(' ', '_')
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@api_router.get("/lesson-plans/{plan_id}/download/powerpoint")
async def download_lesson_plan_powerpoint(plan_id: str, user_id: str = Depends(get_current_user)):
    plan = await db.lesson_plans.find_one({"id": plan_id, "user_id": user_id}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=404, detail="Lesson plan not found")
    
    # Create PowerPoint
    prs = Presentation()
    prs.slide_width = PptxInches(10)
    prs.slide_height = PptxInches(7.5)
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Lesson Plan"
    subtitle.text = f"{plan['subject']} - {plan['topic']}\\n{plan['grade_level']}"
    
    # Content slides
    def add_content_slide(title_text, content_text):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = title_text
        content.text = content_text
    
    add_content_slide("Learning Objectives", plan['objectives'])
    add_content_slide("Materials Needed", plan['materials'])
    add_content_slide("Procedure/Activities", plan['procedure'])
    add_content_slide("Assessment Methods", plan['assessment'])
    
    if plan.get('notes'):
        add_content_slide("Additional Notes", plan['notes'])
    
    # Save to BytesIO
    file_stream = io.BytesIO()
    prs.save(file_stream)
    file_stream.seek(0)
    
    filename = f"{plan['subject']}_{plan['topic']}_lesson_plan.pptx".replace(' ', '_')
    
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# Simple image proxy to work around CORS when inlining remote images client-side
@api_router.get("/proxy/image")
async def proxy_image(url: str):
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            resp = await client.get(url, follow_redirects=True)

        if resp.status_code != 200:
            raise HTTPException(status_code=502, detail="Failed to fetch image")

        content_type = resp.headers.get('content-type', 'application/octet-stream')
        return Response(content=resp.content, media_type=content_type)
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
